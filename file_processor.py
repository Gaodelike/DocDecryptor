import os
from docx import Document
from PyPDF2 import PdfReader, PdfWriter
from PIL import Image
from openpyxl import load_workbook
from pptx import Presentation


def read_and_save_docx(file_path, new_file_path):
    # 读取并保存Word文档
    doc = Document(file_path)
    doc.save(new_file_path)


def read_and_save_pdf(file_path, new_file_path):
    # 读取并保存PDF文档
    reader = PdfReader(file_path)
    writer = PdfWriter()

    for page in reader.pages:
        writer.add_page(page)

    with open(new_file_path, "wb") as output_file:
        writer.write(output_file)


def read_and_save_image(file_path, new_file_path):
    # 读取并保存图像文件
    image = Image.open(file_path)
    image.save(new_file_path)


def read_and_save_xlsx(file_path, new_file_path):
    # 读取并保存Excel文件
    workbook = load_workbook(file_path)
    workbook.save(new_file_path)


def read_and_save_pptx(file_path, new_file_path):
    # 读取并保存PowerPoint文件
    presentation = Presentation(file_path)
    presentation.save(new_file_path)


def is_encrypted(file_path):
    # 判断加密文件是否有对应的ascii字符，若有，则返回True，若无，则返回False
    search_string = "Esafenet".encode("ascii")
    with file_path.open("rb") as file:
        first_chunk = file.read(1024)
        if search_string in first_chunk:
            print(
                f"Found '{search_string.decode()}' at the beginning of file: {file_path}"
            )
            return True
        else:
            print(
                f"'{search_string.decode()}' not found at the beginning of file: {file_path}"
            )
            return False


def process_file(file_path):
    file_extension = os.path.splitext(file_path)[1].lower()
    new_file_path = f"{os.path.splitext(file_path)[0]}.info"

    try:
        if file_extension == '.docx':
            read_and_save_docx(file_path, new_file_path)
        elif file_extension == '.pdf':
            read_and_save_pdf(file_path, new_file_path)
        elif file_extension in ['.jpg', '.jpeg', '.png', '.bmp', '.gif']:
            read_and_save_image(file_path, new_file_path)
        elif file_extension == '.xlsx':
            read_and_save_xlsx(file_path, new_file_path)
        elif file_extension == '.pptx':
            read_and_save_pptx(file_path, new_file_path)
        else:
            if is_encrypted(file_path):
                with open(file_path, "rb") as original_file:
                    binary_data = original_file.read()

                # 使用 .info 扩展名保存加密文件
                with open(new_file_path, "wb") as new_file:
                    new_file.write(binary_data)
                print(f"Encrypted file saved as: {new_file_path}")
            else:
                raise ValueError(f"Unsupported or unencrypted file type: {file_extension}")
    except Exception as e:
        raise e
